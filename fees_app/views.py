from django.shortcuts import render, redirect
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .models import Student, UploadedFile
from openpyxl import load_workbook

# Create your views here.

# Health check endpoint
# This is a lightweight endpoint used by monitoring services (like UptimeRobot)
# to check if the application is running and responsive
# No authentication required - public endpoint
# Does not access database or perform heavy processing
# Returns HTTP 200 OK with simple response
def health_check(request):
    # Return a simple "OK" response
    # This is extremely lightweight and responds quickly
    # Can be called by monitoring services every 5 minutes to keep the service awake
    return HttpResponse("OK", status=200)

# View to handle Excel file upload and display uploaded files
# login_required ensures only authenticated users can access this page
# Redirects to Django admin login if not authenticated
# After login, checks if user is staff or superuser
@login_required(login_url='/admin/login/')
def upload_excel(request):
    # Check if user is staff or superuser (admin access)
    # This check runs after login_required confirms user is authenticated
    if not request.user.is_staff and not request.user.is_superuser:
        # If logged in but not admin, redirect to home page (access denied)
        return redirect('display_students')
    
    # Handle file upload when POST request
    if request.method == 'POST' and 'excel_file' in request.FILES:
        excel_file = request.FILES['excel_file']
        
        # Load the Excel file using openpyxl
        try:
            workbook = load_workbook(excel_file)
            sheet = workbook.active
            
            # Create an UploadedFile record to track this upload
            uploaded_file = UploadedFile.objects.create(
                filename=excel_file.name
            )
            
            # Skip the header row (assuming first row is header)
            # Start from row 2 to read data
            for row in sheet.iter_rows(min_row=2, values_only=True):
                # Extract data from each row
                # Expected columns: Name, Class, Father's Name, Fees Remaining
                name = row[0] if row[0] else ''
                student_class = row[1] if row[1] else ''
                fathers_name = row[2] if row[2] else ''
                fees_remaining = row[3] if row[3] else 0
                
                # Create and save student record associated with this uploaded file
                Student.objects.create(
                    name=name,
                    student_class=student_class,
                    fathers_name=fathers_name,
                    fees_remaining=fees_remaining,
                    uploaded_file=uploaded_file
                )
            
            # Redirect to the same page to show updated file list
            return redirect('upload_excel')
            
        except Exception as e:
            # Handle any errors during file processing
            return HttpResponse(f"Error processing file: {str(e)}")
    
    # Get all uploaded files, ordered by upload date (newest first)
    uploaded_files = UploadedFile.objects.all()
    
    # Render the template with upload form and file list
    return render(request, 'fees_app/upload.html', {
        'uploaded_files': uploaded_files
    })

# View to delete an uploaded file and all its associated students
# login_required ensures only authenticated users can access this page
# After login, checks if user is staff or superuser
@login_required(login_url='/admin/login/')
def delete_file(request, file_id):
    # Check if user is staff or superuser (admin access)
    # This check runs after login_required confirms user is authenticated
    if not request.user.is_staff and not request.user.is_superuser:
        # If logged in but not admin, redirect to home page (access denied)
        return redirect('display_students')
    
    # Only allow POST requests for deletion
    if request.method != 'POST':
        return HttpResponse("Method not allowed", status=405)
    
    try:
        # Get the uploaded file
        uploaded_file = UploadedFile.objects.get(id=file_id)
        
        # Delete the file (this will cascade delete all associated students)
        uploaded_file.delete()
        
        # Redirect back to upload page
        return redirect('upload_excel')
        
    except UploadedFile.DoesNotExist:
        # Handle case where file doesn't exist
        return HttpResponse("File not found", status=404)

# View to display all students with search functionality
# This page remains public - no authentication required
def display_students(request):
    # Get the search query from the request
    search_query = request.GET.get('search', '')
    
    # Filter students based on search query
    if search_query:
        # Search for students whose name contains the search query (case-insensitive)
        students = Student.objects.filter(name__icontains=search_query)
    else:
        # If no search query, show all students
        students = Student.objects.all()
    
    # Render the template with students and search query
    return render(request, 'fees_app/display.html', {
        'students': students,
        'search_query': search_query
    })
